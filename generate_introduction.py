#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new Document
doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Title
title = doc.add_heading('Introduction', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Paragraph 1
p1 = doc.add_paragraph(
    'Triple-negative breast cancer (TNBC), characterized by the absence of estrogen receptor, '
    'progesterone receptor, and HER2 expression, accounts for 15-20% of all breast cancers yet '
    'contributes disproportionately to cancer-related mortality [1, 2]. The lack of targetable '
    'receptors renders TNBC unresponsive to endocrine and HER2-targeted therapies, leaving cytotoxic '
    'chemotherapy as the primary treatment modality with five-year survival rates of 60-77% [3, 4]. '
    'Despite exhibiting the highest levels of tumor-infiltrating lymphocytes (TILs) and PD-L1 '
    'expression among breast cancer subtypes [5, 6], leading to FDA approval of pembrolizumab plus '
    'chemotherapy for PD-L1-positive metastatic TNBC [7, 8], objective response rates remain modest '
    'at 20-40% with frequent development of acquired resistance [9, 10]. This underscores the '
    'critical need to comprehensively dissect immune-tumor interactions within the TNBC '
    'microenvironment, particularly the T cell dynamics that orchestrate anti-tumor immunity.'
)

# Paragraph 2
p2 = doc.add_paragraph(
    'Single-cell RNA sequencing (scRNA-seq) has revolutionized tumor biology by enabling '
    'unprecedented resolution of cellular heterogeneity and mapping of intercellular communication '
    'networks [11, 12]. Recent scRNA-seq studies in TNBC have revealed immunologically distinct '
    'subtypes with differential therapeutic responses [13, 14], yet comprehensive characterization '
    'of T cell composition, functional states, and their regulatory mechanisms remains incomplete. '
    'T cell function is precisely controlled through post-translational modifications, with '
    'ubiquitination playing pivotal roles in TCR signaling, activation thresholds, and immune '
    'checkpoint regulation [15, 16]. The ubiquitin-proteasome system regulates diverse cellular '
    'processes through coordinated action of E3 ligases and deubiquitinases [17, 18]. For instance, '
    'the E3 ligase Cbl-b negatively regulates TCR signaling [19], while the deubiquitinase OTUB1 '
    'stabilizes PD-L1 to promote immune evasion [20, 21]. Despite ubiquitination\'s established role '
    'in cancer progression and immune regulation [22, 23], the contribution of T cell-related '
    'ubiquitination genes to TNBC pathogenesis and their prognostic implications remain unexplored.'
)

# Paragraph 3
p3 = doc.add_paragraph(
    'Current TNBC prognostic models rely primarily on clinicopathological parameters with limited '
    'accuracy [24, 25]. While PAM50 and Oncotype DX have improved risk stratification in '
    'hormone receptor-positive breast cancer, they show variable performance in TNBC [26, 27]. '
    'Recent TNBC-specific gene signatures targeting immune profiles, metabolism, or epithelial-mesenchymal '
    'transition (EMT) [28-31] often lack mechanistic grounding and fail to integrate the dynamics of '
    'immune-tumor crosstalk. Furthermore, bulk transcriptomics-derived signatures cannot capture the '
    'cellular heterogeneity and communication networks that fundamentally shape tumor biology [32, 33]. '
    'Despite evidence linking T cell infiltration to TNBC outcomes [5, 34], three critical gaps '
    'persist: (1) comprehensive single-cell characterization of T cell composition and transcriptional '
    'states is lacking; (2) ubiquitination-mediated T cell regulation in TNBC and prognostic '
    'correlations with specific ubiquitination genes remain unknown; (3) no prognostic models integrate '
    'T cell molecular features with ubiquitination pathway information for accurate risk stratification.'
)

# Paragraph 4
p4 = doc.add_paragraph(
    'To address these gaps, we integrated single-cell and bulk transcriptomic analyses of TNBC versus '
    'normal breast tissue. We performed scRNA-seq on 16 TNBC tumor samples and 3 normal tissues '
    '(84,837 cells total), characterized the cellular landscape, and mapped intercellular communication '
    'networks using CellChat [35]. We identified T cell-related ubiquitination genes (TCRUG) by '
    'intersecting T cell differentially expressed genes with 807 curated ubiquitination-related genes [36]. '
    'Through weighted gene co-expression network analysis (WGCNA) [37], differential expression '
    'analysis, survival correlation, and LASSO-Cox regression [38], we developed and validated a '
    'three-gene prognostic signature comprising GMCL1, KRT8, and OTUB1 across multiple independent '
    'cohorts (TCGA-TNBC training n=82, internal test n=34, external validation GSE25066 n=508). '
    'We further constructed an integrated clinical-molecular nomogram for practical application. '
    'Our comprehensive approach reveals the interplay between T cells and ubiquitination in TNBC, '
    'identifies MIF and CXCL10-CXCR3 signaling as potential therapeutic targets, and provides a '
    'mechanistically grounded prognostic tool for TNBC patient stratification.'
)

# Save the document
doc.save('/home/user/data_bricks/introduc.docx')
print("Document generated successfully: introduc.docx")
